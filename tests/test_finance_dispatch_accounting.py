"""Finance must integrate MW once and include firming costs in cashflows."""
from dataclasses import replace
from io import BytesIO

import numpy as np
import openpyxl
import pandas as pd
import pytest

from ppa.financial_model import (
    EnergyInputs, ProjectFinanceInputs, energy_inputs_from_result,
    run_project_finance,
)
from ppa.financial_model_excel import export_financial_model, export_hourly_timeseries
from ppa.financials import run_multi_year_financial_analysis
from ppa.results import DispatchSeries, OptimisationResult, RevenueBreakdown, SummaryVolumes
from ppa.scenario import Scenario


def _dispatch_result(minutes=60):
    idx = pd.date_range("2025-01-01", periods=24 * 60 // minutes, freq=f"{minutes}min")
    dt = minutes / 60
    def series(value):
        return pd.Series(value, index=idx, dtype=float)
    d = DispatchSeries(
        wind_gen=series(10), pv_gen=series(5), market_buy=series(2),
        allowed_shortfall=series(0), penalty_gen=series(1), market_sell=series(3),
        bess_dispatch=series(4), bess_store=series(4), soc=series(1),
        ppa_delivery=series(12),
    )
    s = Scenario(
        name="Accounting test", market_spread=3, transmission_cost_aud_mwh=4,
        simulation_years=25, dispatch_mode="representative_year",
        price_escalation_rate=0.05, lgc_price_aud_mwh=6,
    )
    summ = SummaryVolumes(
        total_load_mwh=312, ppa_delivered_mwh=288,
        renewable_and_storage_to_ppa_mwh=240, market_buy_to_ppa_mwh=48,
        allowed_shortfall_mwh=0, penalty_mwh=24, sold_to_market_mwh=72,
        wind_generation_mwh=240, pv_generation_mwh=120,
        bess_dispatch_mwh=96, bess_charge_mwh=96,
        fulfilled_share=12/13, allowed_shortfall_share_actual=0,
        buy_share_of_ppa_delivery=1/6, penalty_share_of_load=1/13,
    )
    price = series(100)
    rev = RevenueBreakdown(
        ppa_revenue=288*s.ppa_price, excess_revenue=72*(100+6-3),
        market_purchase_cost=48*(100+3+6), penalty_cost=24*s.penalty_price,
        transmission_cost=288*4,
        net_revenue=288*s.ppa_price+72*103-48*109-24*s.penalty_price-288*4,
        effective_capture_price=0,
    )
    return OptimisationResult(s, d, summ, rev, "ok", "optimal", len(idx)*dt, price)


@pytest.mark.parametrize("minutes", [60, 30, 15, 5])
def test_energy_inputs_are_resolution_invariant_and_exclude_storage(minutes):
    result = _dispatch_result(minutes)
    energy = energy_inputs_from_result(result)
    scale = 8760 / 24
    assert energy.ppa_gwh == pytest.approx(288 * scale / 1000)
    assert energy.penalty_gwh == pytest.approx(24 * scale / 1000)
    assert energy.excess_solar_gwh + energy.excess_nonsolar_gwh == pytest.approx(72 * scale / 1000)
    assert energy.marketbuy_gwh == pytest.approx(48 * scale / 1000)
    assert energy.total_solar_gwh + energy.total_nonsolar_gwh == pytest.approx(360 * scale / 1000)
    assert energy.purchase_price == 109
    assert energy.sell_solar_price == 97
    assert energy.sell_nonsolar_price == 97
    assert energy.transmission_cost_aud == 288 * 4 * scale
    unscaled = energy_inputs_from_result(result, annualise=False)
    assert unscaled.marketbuy_gwh == pytest.approx(0.048)
    assert unscaled.total_solar_gwh + unscaled.total_nonsolar_gwh == pytest.approx(0.360)


def _energy():
    return EnergyInputs(
        onsw_mw=100, pv_mw=0, bess_mw=0, bess_mwh=0, load_mw=50,
        ppa_gwh=300, excess_solar_gwh=20, excess_nonsolar_gwh=30,
        penalty_gwh=0, total_solar_gwh=100, total_nonsolar_gwh=250,
        sell_solar_price=60, sell_nonsolar_price=80,
    )


@pytest.mark.parametrize("horizon", [1, 17, 25, 30])
def test_fast_financial_horizon_controls_cashflows_and_annuity(horizon):
    result = _dispatch_result()
    scenario = replace(result.scenario, simulation_years=horizon, project_life_yrs=25)
    projected = run_multi_year_financial_analysis(scenario, [result])
    aligned = run_multi_year_financial_analysis(
        replace(scenario, project_life_yrs=horizon), [result]
    )
    assert len(projected.yearly) == horizon
    assert len(projected.cumulative_npv) == horizon
    assert projected.npv == pytest.approx(aligned.npv)
    assert projected.lcoe == pytest.approx(aligned.lcoe)


def test_firming_costs_reduce_cashflows_tax_and_debt():
    p = ProjectFinanceInputs(ancillary_pct=0, max_gearing_contracted=1)
    e = _energy()
    base = run_project_finance(p, e)
    costs = run_project_finance(p, replace(e, marketbuy_gwh=50, purchase_price=100,
                                          transmission_cost_aud=1_000_000))
    first_ops = int(np.flatnonzero(costs.schedule["ops_flag"])[0])
    for i in np.flatnonzero(costs.schedule["ppa_flag"]):
        expected = 5 * (1+p.nonsolar_price_inflation)**i + (1+p.cost_inflation)**i
        assert costs.schedule["ebitda"][i] == pytest.approx(base.schedule["ebitda"][i]-expected)
    assert costs.schedule["market_purchase_cost"][first_ops] > 5
    assert costs.schedule["transmission_cost"][first_ops] > 1
    assert costs.schedule["tax"].sum() < base.schedule["tax"].sum()
    assert costs.total_debt < base.total_debt
    assert costs.schedule["fcff"].sum() < base.schedule["fcff"].sum()
    assert costs.schedule["fcfe"].sum() < base.schedule["fcfe"].sum()
    assert costs.project_irr < base.project_irr
    assert costs.equity_irr < base.equity_irr
    post_ppa = costs.schedule["ops_flag"].astype(bool) & ~costs.schedule["ppa_flag"].astype(bool)
    assert np.all(costs.schedule["market_purchase_cost"][post_ppa] == 0)
    assert np.all(costs.schedule["transmission_cost"][post_ppa] == 0)


def test_actual_year_costs_are_not_price_escalated_twice():
    e = replace(_energy(), marketbuy_gwh=10, purchase_price=100, transmission_cost_aud=200_000)
    p = ProjectFinanceInputs(escalate_merchant_prices=False, cost_inflation=0)
    actual = run_project_finance(p, e, annual_energy=[e, replace(e, purchase_price=120)])
    ops = np.flatnonzero(actual.schedule["ops_flag"])
    assert actual.schedule["market_purchase_cost"][ops[0]] == 1
    assert actual.schedule["market_purchase_cost"][ops[1]] == 1.2
    last_ppa = np.flatnonzero(actual.schedule["ppa_flag"])[-1]
    assert actual.schedule["market_purchase_cost"][last_ppa] == 1.2
    assert actual.schedule["transmission_cost"][ops[0]] == 0.2


def test_negative_cfads_never_creates_negative_debt():
    e = replace(_energy(), marketbuy_gwh=1000, purchase_price=1000)
    result = run_project_finance(ProjectFinanceInputs(), e)
    assert result.total_debt == 0
    assert np.all(result.schedule["debt_draw"] >= 0)
    assert np.all(result.schedule["debt_closing"] >= -1e-8)


def test_single_dispatch_projects_25_compact_financial_years():
    result = _dispatch_result()
    financials = run_multi_year_financial_analysis(result.scenario, [result], 2025)
    assert len(financials.yearly) == 25
    assert financials.yearly[-1].year == 2049
    assert financials.total_lifetime_generation_mwh == 25 * 360
    factor = 1.05**24
    last = financials.yearly[-1]
    assert last.ppa_revenue == result.revenue.ppa_revenue
    assert last.transmission_cost == result.revenue.transmission_cost
    assert last.merch_revenue == pytest.approx(72*(100*factor+6-3))
    assert last.market_buy_cost == pytest.approx(48*(100*factor+3+6))
    assert last.net_cashflow == pytest.approx(last.net_revenue-financials.annual_opex)
    assert len(result.dispatch.wind_gen) == 24


def test_subhourly_excel_rows_use_mwh_and_exclude_battery_generation():
    result = _dispatch_result(minutes=30)
    sheet = openpyxl.load_workbook(BytesIO(export_hourly_timeseries([result])))["Hourly"]
    assert sheet.cell(17, 4).value == 5  # 10 MW over half an hour
    assert sheet.cell(17, 8).value == 7.5  # wind + PV, excluding BESS
    assert sheet.cell(17, 9).value == 1
    assert sheet.cell(17, 15).value == 97
    assert sheet.cell(17, 16).value == 109
    assert sheet.cell(17, 17).value == 24
    assert "*0.5" in sheet.cell(4, 3).value


def test_excel_deducts_costs_in_live_cashflow_chain():
    p = ProjectFinanceInputs()
    e = replace(_energy(), marketbuy_gwh=50, purchase_price=100, transmission_cost_aud=1_000_000)
    sheet = openpyxl.load_workbook(BytesIO(export_financial_model(p, e, run_project_finance(p,e))))["Model"]
    labels = {sheet.cell(i, 2).value: i for i in range(1, sheet.max_row+1)}
    net = sheet.cell(labels["Net contracted revenue"], 4).value
    assert f"-D{labels['Market purchase cost']}" in net
    assert f"-D{labels['Transmission cost']}" in net


def test_excel_preserves_actual_year_volumes_and_purchase_prices():
    e = replace(_energy(), marketbuy_gwh=10, purchase_price=100)
    second = replace(e, ppa_gwh=280, purchase_price=120)
    p = ProjectFinanceInputs(escalate_merchant_prices=False)
    result = run_project_finance(p, e, annual_energy=[e, second])
    wb = openpyxl.load_workbook(BytesIO(export_financial_model(p,e,result)))
    annual = wb["Annual Energy"]
    rows = {annual.cell(i,1).value:i for i in range(2,annual.max_row+1)}
    assert annual.cell(rows["purchase_price"],2).value == 100
    assert annual.cell(rows["purchase_price"],3).value == 120
    assert annual.cell(rows["ppa_gwh"],3).value == 280
    model = wb["Model"]
    labels = {model.cell(i,2).value:i for i in range(1,model.max_row+1)}
    ops = np.flatnonzero(result.schedule["ops_flag"])
    first = model.cell(labels["Average market purchase price"],4+int(ops[0])).value
    second_formula = model.cell(labels["Average market purchase price"],4+int(ops[1])).value
    assert "$B$" in first
    assert "$C$" in second_formula
    assert "'Annual Energy'!" in first
