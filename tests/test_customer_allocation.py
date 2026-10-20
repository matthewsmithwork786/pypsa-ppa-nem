from dataclasses import replace
from io import BytesIO

import numpy as np
import openpyxl
import pandas as pd
import pytest

from ppa.financial_model import (energy_inputs_from_result, energy_inputs_from_results,
                                 ProjectFinanceInputs, run_project_finance)
from ppa.financial_model_excel import export_hourly_timeseries, export_financial_model
from ppa.network import build_network
from ppa.results import DispatchSeries, customer_first_dispatch, extract_results, merchant_sales_frame
from ppa.scenario import Scenario
from ppa.solver import solve


def _dispatch(index, load, delivery=None, sales=None):
    zero = pd.Series(0., index=index)
    delivery = pd.Series(delivery if delivery is not None else 0., index=index)
    return DispatchSeries(zero, zero, zero, load - delivery, zero,
                          pd.Series(sales if sales is not None else 0., index=index),
                          zero, zero, zero, delivery)


@pytest.mark.parametrize('dt', [0.5, 1.0])
@pytest.mark.parametrize('battery', [False, True])
def test_allocation_preserves_solve_and_reconciles_all_downstream_results(dt, battery):
    s = Scenario(onsw_mw=150, pv_mw=0, ppaload_mw=100, include_bess=battery,
                 bess_mw=20, bess_mwh=40, required_delivery_share=.75,
                 lgc_price_aud_mwh=5, transmission_cost_aud_mwh=3)
    ts = pd.DataFrame({'ts_PVGen': 0., 'ts_WindGen': [1., .3, 0., .8] * 6,
                       'ts_MktPrice': [2000., 20., 100., 1500.] * 6,
                       'ppaload_mw': [100., 100., 100., 50.] * 6},
                      index=pd.date_range('2025-01-01', periods=24,
                                          freq=pd.Timedelta(hours=dt)))
    n = build_network(ts, s, resolution_h=dt)
    status, condition = solve(n, s, ts)
    assert condition == 'optimal'
    raw = extract_results(n, s, ts, status, condition, resolution_h=dt)
    allocated_s = replace(s, customer_first_allocation=True)
    r = extract_results(n, allocated_s, ts, status, condition, resolution_h=dt)
    d = r.dispatch
    assert raw.raw_dispatch is None
    assert r.raw_dispatch is not None
    for field in ['wind_gen', 'pv_gen', 'market_buy', 'bess_dispatch', 'bess_store', 'soc']:
        pd.testing.assert_series_equal(getattr(d, field), getattr(raw.dispatch, field))
    transfer = np.minimum(raw.dispatch.market_sell.clip(lower=0),
                          (ts.ppaload_mw - raw.dispatch.ppa_delivery).clip(lower=0))
    assert transfer.sum() > 1
    np.testing.assert_allclose(d.ppa_delivery, raw.dispatch.ppa_delivery + transfer)
    np.testing.assert_allclose(d.market_sell, raw.dispatch.market_sell - transfer, atol=1e-6)
    np.testing.assert_allclose(d.ppa_delivery + d.allowed_shortfall + d.penalty_gen, ts.ppaload_mw)
    assert not ((d.market_sell > 1e-6) & ((ts.ppaload_mw - d.ppa_delivery) > 1e-6)).any()
    gap_mwh = float((ts.ppaload_mw - d.ppa_delivery).sum()) * dt
    expected_penalty = max(0, gap_mwh - ts.ppaload_mw.sum() * dt * s.allowed_shortfall_share)
    assert r.summary.penalty_mwh == pytest.approx(expected_penalty, abs=1e-6)
    assert r.revenue.ppa_revenue == pytest.approx(d.ppa_delivery.sum() * dt * s.ppa_price)
    assert r.revenue.penalty_cost == pytest.approx(expected_penalty * s.penalty_price, abs=1e-6)
    assert r.revenue.excess_revenue == pytest.approx(
        (d.market_sell * (ts.ts_MktPrice + s.lgc_price_aud_mwh - s.market_spread)).sum() * dt)
    assert merchant_sales_frame(r)['Merchant revenue (A$)'].sum() == pytest.approx(r.revenue.excess_revenue)
    energy = energy_inputs_from_result(r, annualise=False)
    assert energy.ppa_gwh * 1000 == pytest.approx(r.summary.ppa_delivered_mwh)
    assert (energy.excess_solar_gwh + energy.excess_nonsolar_gwh) * 1000 == pytest.approx(r.summary.sold_to_market_mwh)
    assert energy.penalty_gwh * 1000 == pytest.approx(expected_penalty, abs=1e-6)
    assert energy_inputs_from_results([r]).customer_first_allocation
    # Original network and the off-mode result retain the solved settlement.
    again = extract_results(n, s, ts, status, condition, resolution_h=dt)
    pd.testing.assert_series_equal(again.dispatch.market_sell, raw.dispatch.market_sell)
    assert again.revenue == raw.revenue
    np.testing.assert_allclose(
        d.wind_gen + d.pv_gen + d.bess_dispatch - d.bess_store + d.market_buy,
        d.ppa_delivery + d.market_sell, atol=1e-6)
    if dt == 1 and not battery:
        from ppa.financials import run_multi_year_financial_analysis
        from ppa.counterfactuals import compute_counterfactuals

        financials = run_multi_year_financial_analysis(allocated_s, [r])
        assert financials.yearly[0].net_revenue == pytest.approx(r.revenue.net_revenue)
        cf = compute_counterfactuals(ts, allocated_s, r, dt=dt)
        expected_bill = (s.ppa_price * d.ppa_delivery
                         + ts.ts_MktPrice * (ts.ppaload_mw - d.ppa_delivery)).sum() * dt
        assert cf.ppa_offtaker_cost == pytest.approx(expected_bill)
        p = ProjectFinanceInputs()
        fm = run_project_finance(p, energy)
        financial_wb = openpyxl.load_workbook(BytesIO(export_financial_model(p, energy, fm)))
        assert 'Customer-first allocation' in financial_wb['Energy']['B3'].value
    wb = openpyxl.load_workbook(BytesIO(export_hourly_timeseries([r])), data_only=False)
    ws = wb['Hourly']
    from ppa.financial_model_excel import _HOURLY_DATA_START
    rows = range(_HOURLY_DATA_START, _HOURLY_DATA_START + len(ts))
    assert sum(ws.cell(i, 10).value for i in rows) == pytest.approx(r.summary.sold_to_market_mwh, abs=.024)
    assert sum(ws.cell(i, 11).value for i in rows) == pytest.approx(r.summary.ppa_delivered_mwh, abs=.024)


def test_allowance_is_reused_after_allocation_and_resets_each_year():
    idx = pd.to_datetime(['2025-01-01', '2026-01-01'])
    load = pd.Series([100., 100.], index=idx)
    raw = _dispatch(idx, load, sales=[100., 0.])
    d = customer_first_dispatch(raw, load, Scenario(required_delivery_share=.75))
    assert list(d.ppa_delivery) == [100, 0]
    assert list(d.allowed_shortfall) == [0, 25]
    assert list(d.penalty_gen) == [0, 75]
    assert list(raw.market_sell) == [100, 0]


def test_nested_allowance_caps_and_no_allowance():
    idx = pd.to_datetime(['2025-01-01', '2025-01-02', '2025-02-01'])
    load = pd.Series(100., index=idx)
    raw = _dispatch(idx, load, delivery=[10., 60., 100.])
    s = Scenario(required_delivery_share=.95, sla_monthly_enabled=True,
                 sla_monthly_share=.9, sla_daily_enabled=True, sla_daily_share=.8)
    d = customer_first_dispatch(raw, load, s)
    assert list(d.allowed_shortfall) == pytest.approx([7.5, 7.5, 0])
    no_allowance = customer_first_dispatch(raw, load, replace(s, enable_shortfall=False))
    assert no_allowance.allowed_shortfall.sum() == 0
    assert no_allowance.penalty_gen.sum() == 130


def test_zero_load_and_no_exports_are_safe():
    idx = pd.date_range('2025-01-01', periods=3, freq='h')
    load = pd.Series([0., 100., 100.], index=idx)
    raw = _dispatch(idx, load, delivery=[0., 50., 100.], sales=[10., 0., 0.])
    d = customer_first_dispatch(raw, load, Scenario())
    pd.testing.assert_series_equal(d.ppa_delivery, raw.ppa_delivery)
    pd.testing.assert_series_equal(d.market_sell, raw.market_sell)
    assert np.isfinite(d.allowed_shortfall).all()


def test_setting_defaults_off_round_trips_and_is_editable_in_ui():
    from streamlit.testing.v1 import AppTest

    assert not Scenario().customer_first_allocation
    assert Scenario(**replace(Scenario(), customer_first_allocation=True).to_dict()).customer_first_allocation
    app = AppTest.from_string("""
import streamlit as st
from ppa.scenario import Scenario
from ui.scenario_form import render_scenario_form
s = render_scenario_form(Scenario())
st.session_state['chosen_allocation'] = s.customer_first_allocation
""").run(timeout=20)
    assert not app.exception
    assert not app.session_state['chosen_allocation']
    app.toggle(key='sf_customer_first_allocation').set_value(True).run()
    assert not app.exception
    assert app.session_state['chosen_allocation']


def test_resetting_scenario_resets_allocation_widget_and_export_cache(monkeypatch):
    from ui import state

    session = {'sf_customer_first_allocation': True, 'fm_hourly_xlsx_for': 'old',
               'fm_hourly_xlsx': b'old workbook'}
    monkeypatch.setattr(state.st, 'session_state', session)
    state.set_scenario(Scenario())
    assert 'sf_customer_first_allocation' not in session
    assert 'fm_hourly_xlsx' not in session
    assert not session[state.SCENARIO_KEY].customer_first_allocation
