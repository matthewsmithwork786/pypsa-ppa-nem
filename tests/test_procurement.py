from dataclasses import replace

import numpy as np
import pandas as pd
import pytest

from ppa.procurement import Procurement, optimise_procurement, validate_inputs


def year_data(wind=0.4):
    index = pd.date_range("2025-01-01", periods=8760, freq="h")
    return pd.DataFrame({"ts_PVGen": np.where((index.hour >= 6) & (index.hour < 18), 0.8, 0),
                         "ts_WindGen": wind, "ts_LoadMW": 90.0,
                         "ts_MktPrice": np.where(index.hour < 6, -20.0, 150.0)}, index=index)


@pytest.fixture(scope="module")
def nsw_result():
    return optimise_procurement(year_data(wind=0.37), Procurement(peak_mw=101))


@pytest.fixture(scope="module")
def aemo_result():
    return optimise_procurement(year_data(wind=0), Procurement(mode="AEMO"))


def test_nsw_energy_mix_storage_and_annual_matching(nsw_result):
    r = nsw_result
    d = r["dispatch"]
    assert d.Wind.sum() >= 0.4 * d.Load.sum() - 1e-4
    assert d.PV.sum() >= 0.6 * d.Load.sum() - 1e-4
    assert abs(d.Wind.sum() / (d.Wind.sum() + d.PV.sum()) - 0.4) < 0.01
    assert all(float(v).is_integer() for v in r["capacities"].values())
    assert (d.Wind + d.PV - d.Curtailed).sum() >= d.Load.sum() - 1e-4
    assert r["capacities"]["BESS"] == 26
    assert r["fees"]["BESS fee"] == pytest.approx(26 * 4 * 6000 * 12)


def test_aemo_nighttime_targets_require_storage_without_wind(aemo_result):
    r = aemo_result
    assert r["capacities"]["Wind"] == pytest.approx(0)
    assert r["capacities"]["BESS"] > 0
    assert (r["compliance"].Achieved >= r["compliance"].Target - 1e-6).all()
    d = r["dispatch"]
    # Charging cannot originate from the grid, and the year is energy-neutral in SOC.
    assert (d["Battery charge"] <= d.PV + d.Wind + 1e-4).all()
    assert d["Battery discharge"].sum() == pytest.approx(d["Battery charge"].sum() * 0.95**2)


@pytest.mark.parametrize("fixture", ["nsw_result", "aemo_result"])
def test_energy_balances_and_customer_cashflows(fixture, request):
    r = request.getfixturevalue(fixture)
    d, fees = r["dispatch"], r["fees"]
    assert not ((d["Battery charge"] > 1e-5) & (d["Battery discharge"] > 1e-5)).any()
    np.testing.assert_allclose(d["Renewable delivery"] + d["Spot purchases"], d.Load, atol=1e-5)
    np.testing.assert_allclose(d.PV + d.Wind + d["Battery discharge"],
                               d["Renewable delivery"] + d["Battery charge"] + d["Excess sold"] + d.Curtailed, atol=1e-5)
    assert r["spot_cost"] == pytest.approx((d.Load * d["Spot price"]).sum())
    assert fees["Spot shortfalls"] == pytest.approx((d["Spot purchases"] * d["Spot price"]).sum())
    assert fees["Excess revenue"] == pytest.approx((d["Excess sold"] * d["Spot price"]).sum())
    expected = fees["PV PPA"] + fees["Wind PPA"] + fees["BESS fee"] + fees["Spot shortfalls"] - fees["Excess revenue"]
    assert r["net_cost"] == pytest.approx(expected)
    assert r["savings"] == pytest.approx(r["spot_cost"] - expected)


def test_infeasible_wind_share_is_reported():
    with pytest.raises(ValueError, match="No optimal portfolio"):
        optimise_procurement(year_data(wind=0), Procurement())


@pytest.mark.parametrize("problem", ["gap", "nan", "infinity", "cf", "peak", "duplicate"])
def test_invalid_uploads_are_blocked(problem):
    ts = year_data()
    if problem == "gap":
        ts = ts.iloc[1:]
    elif problem in {"nan", "infinity"}:
        ts.iloc[0, 0] = np.nan if problem == "nan" else np.inf
    elif problem == "cf":
        ts.iloc[0, 0] = 1.1
    elif problem == "peak":
        ts.iloc[0, 2] = 101
    else:
        ts = pd.concat([ts, ts.iloc[-1:]])
    with pytest.raises(ValueError):
        validate_inputs(ts, Procurement())


def test_leap_year_is_supported():
    ts = year_data().reindex(pd.date_range("2024-01-01", periods=8784, freq="h")).fillna(
        {"ts_PVGen": 0.2, "ts_WindGen": 0.4, "ts_LoadMW": 90, "ts_MktPrice": 80})
    validate_inputs(ts, Procurement())


def test_bad_assumptions_are_blocked():
    with pytest.raises(ValueError):
        validate_inputs(year_data(), replace(Procurement(), band_shares=(0.2, 0.4)))


def test_chosen_capacity_has_no_matching_targets_and_reports_all_metrics():
    ts = year_data()
    ts["ts_MktPrice"] = 100.0
    s = Procurement(mode="Fixed", pv_mw=0, wind_mw=100, bess_mw=0, bess_mwh=0,
                    annual_match=3, wind_share=0.99, band_shares=(1, 1, 1, 1))
    r = optimise_procurement(ts, s)
    assert r["capacities"] == {"PV": 0, "Wind": 100, "BESS": 0}
    assert r["compliance"].empty
    metrics = r["matching"].set_index("Metric").Achieved
    assert len(metrics) == 8
    assert metrics["BESS power / peak load"] == 0
    for label, value in metrics.items():
        if label != "BESS power / peak load":
            assert value == pytest.approx(40 / 90)


def test_chosen_bess_power_and_energy_are_independent():
    ts = year_data().astype("float32")
    r = optimise_procurement(ts, Procurement(mode="Fixed", pv_mw=180, wind_mw=75, bess_mw=13, bess_mwh=78))
    assert r["capacities"] == {"PV": 180, "Wind": 75, "BESS": 13}
    assert r["assumptions"].storage_duration_h == 6
    assert r["fees"]["BESS fee"] == pytest.approx(78 * 6000 * 12)
    d = r["dispatch"]
    assert d["Stored energy"].max() <= 78 + 1e-5
    assert d["Battery discharge"].max() <= 13 + 1e-5
    assert not ((d["Battery charge"] > 1e-5) & (d["Battery discharge"] > 1e-5)).any()
    fees = r["fees"]
    expected = fees["PV PPA"] + fees["Wind PPA"] + fees["BESS fee"] + fees["Spot shortfalls"] - fees["Excess revenue"]
    assert r["net_cost"] == pytest.approx(expected, rel=0, abs=1e-5)


def test_time_matching_weights_load_and_excludes_surplus():
    ts = year_data()
    daytime = (ts.index.hour >= 6) & (ts.index.hour < 18)
    ts["ts_LoadMW"] = np.where(daytime, 90.0, 10.0)
    ts["ts_MktPrice"] = 100.0
    r = optimise_procurement(ts, Procurement(mode="Fixed", pv_mw=0, wind_mw=100, bess_mw=0, bess_mwh=0))
    metrics = r["matching"].set_index("Metric").Achieved
    assert metrics["Annual renewable volume matching"] == pytest.approx(0.8)
    assert metrics["Wind generation / annual load"] == pytest.approx(0.8)
    assert metrics["Time matching · whole year"] == pytest.approx(0.5)
    assert metrics["Time matching · 00:00–06:00"] == pytest.approx(1)
    assert metrics["Time matching · 06:00–12:00"] == pytest.approx(40 / 90)


def test_zero_chosen_capacities_can_buy_all_load_at_spot():
    r = optimise_procurement(year_data(), Procurement(mode="Fixed", pv_mw=0, wind_mw=0, bess_mw=0, bess_mwh=0))
    assert r["net_cost"] == pytest.approx(r["spot_cost"])
    assert (r["matching"].Achieved == 0).all()


@pytest.mark.parametrize("values", [dict(pv_mw=-1), dict(pv_mw=10.5), dict(wind_mw=10.5), dict(bess_mw=10.5), dict(bess_mw=0, bess_mwh=10), dict(bess_mw=10, bess_mwh=0)])
def test_invalid_fixed_capacities_are_blocked(values):
    with pytest.raises(ValueError):
        validate_inputs(year_data(), replace(Procurement(mode="Fixed"), **values))


def test_chosen_generation_requires_a_profile():
    with pytest.raises(ValueError, match="Wind profile"):
        validate_inputs(year_data(wind=0), Procurement(mode="Fixed"))


def test_time_matching_is_energy_weighted_and_excludes_exports(nsw_result):
    r = nsw_result
    d = r["dispatch"]
    metrics = r["matching"].set_index("Metric").Achieved
    assert metrics["Time matching · whole year"] == pytest.approx(d["Renewable delivery"].sum() / d.Load.sum())
    assert metrics["Wind generation / annual load"] == pytest.approx(d.Wind.sum() / d.Load.sum())
    assert metrics["Annual renewable volume matching"] > metrics["Time matching · whole year"]


def test_aemo_capacities_are_whole_mw(aemo_result):
    assert all(float(v).is_integer() for v in aemo_result["capacities"].values())


def test_source_allocations_and_cost_bridge_reconcile(nsw_result):
    from ppa.procurement_reporting import cost_breakdown, SOURCES
    d = nsw_result["dispatch"]
    np.testing.assert_allclose(d[[f"{s} excess" for s in SOURCES]].sum(axis=1), d["Excess sold"], atol=1e-5)
    np.testing.assert_allclose(d[[f"{s} to load" for s in SOURCES]].sum(axis=1), d["Renewable delivery"], atol=1e-5)
    np.testing.assert_allclose(d["PV charged"] + d["Wind charged"], d["Battery charge"], atol=1e-5)
    assert (d["BESS excess"] == 0).all()
    np.testing.assert_allclose(d["BESS to load"], d["Battery discharge"], atol=1e-5)
    bridge = cost_breakdown(nsw_result)
    assert bridge["Annual cost (A$)"].sum() == pytest.approx(nsw_result["net_cost"])
    for row in bridge.itertuples(index=False, name=None):
        label, qty, unit, rate, _, cost = row
        factor = 12 if unit == "MWh capacity" else 1000
        assert cost == pytest.approx(qty * factor * rate * (-1 if "excess" in label else 1))


def test_workbook_formulas_cached_values_format_and_sheets(nsw_result):
    from io import BytesIO
    from openpyxl import load_workbook
    from ppa.procurement_excel import export_procurement
    payload = export_procurement(nsw_result)
    wb = load_workbook(BytesIO(payload), data_only=False)
    values = load_workbook(BytesIO(payload), data_only=True)
    assert wb.sheetnames == ["Annual financials", "Hourly dispatch"]
    assert wb["Hourly dispatch"].max_row == 8761
    assert wb["Hourly dispatch"].freeze_panes == "B2"
    ws = wb["Annual financials"]
    cached = values["Annual financials"]
    assert ws["E15"].value == "=E13+E14"
    assert cached["E15"].value == pytest.approx(nsw_result["net_cost"])
    assert sum(cached[f"E{i}"].value for i in range(22, 29)) == pytest.approx(nsw_result["net_cost"])
    assert cached["E16"].value == pytest.approx(nsw_result["spot_cost"])
    assert cached["E17"].value == pytest.approx(nsw_result["savings"])
    assert cached["E18"].value == pytest.approx(nsw_result["net_cost"] / nsw_result["dispatch"].Load.sum())
    assert ws["E15"].font.bold
    assert "0.00" in ws["E15"].number_format
    assert ws["D7"].value == 6000
    assert ws["E24"].value == "=B24*12*D24"
    assert ws["E26"].value == "=-B26*1000*D26"
    assert ws["D25"].value.startswith("=IF(SUM(")
